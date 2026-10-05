from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import unicodedata

from openpyxl import load_workbook

from .models import Product


HEADER_ALIASES = {
    "sku": ("sku", "codigo", "codigo interno", "codigo de barras"),
    "name": ("nombre", "producto", "descripcion"),
    "commercial_name": ("nombre comercial", "nombre profesional", "nombre de venta"),
    "category": ("categoria", "rubro"),
    "quality": ("calidad", "linea", "calidad o linea"),
    "brand": ("marca",),
    "supplier": ("proveedor",),
    "sale_unit": ("se vende por", "unidad de venta", "unidad"),
    "package_weight_kg": ("peso de bolsa kg", "peso bolsa", "peso kg"),
    "price": ("precio", "precio habitual"),
    "promotional_price": ("precio promocional", "precio promo"),
    "promotion_starts_on": ("promocion desde", "promo desde"),
    "promotion_ends_on": ("promocion hasta", "promo hasta"),
    "stock": ("stock", "stock actual"),
    "minimum_stock": ("stock minimo", "minimo", "stock min"),
    "batch": ("lote",),
    "expires_on": ("vencimiento", "vence", "fecha vencimiento"),
    "requires_prescription": ("requiere receta", "receta"),
}


def normalize(value):
    value = "" if value is None else str(value).strip().lower()
    return " ".join("".join(char for char in unicodedata.normalize("NFD", value) if unicodedata.category(char) != "Mn").split())


def as_text(value):
    return "" if value is None else str(value).strip()


def as_decimal(value, field, required=False, default=None):
    if value in (None, ""):
        if required:
            raise ValueError(f"{field} es obligatorio.")
        return default
    text = as_text(value).replace("$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as error:
        raise ValueError(f"{field} debe ser numérico.") from error
    if result < 0:
        raise ValueError(f"{field} no puede ser negativo.")
    return result


def as_date(value, field):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(as_text(value), pattern).date()
        except ValueError:
            continue
    raise ValueError(f"{field} debe tener formato dd/mm/aaaa o aaaa-mm-dd.")


def as_bool(value):
    return normalize(value) in {"si", "sí", "true", "1", "x", "yes"}


def choice_value(value, choices, field, aliases=None, default=None):
    normalized = normalize(value)
    if not normalized and default is not None:
        return default
    mapping = {normalize(key): key for key, _ in choices}
    mapping.update({normalize(label): key for key, label in choices})
    mapping.update(aliases or {})
    if normalized not in mapping:
        raise ValueError(f"{field} no es válido: {as_text(value)}.")
    return mapping[normalized]


def parse_product_workbook(uploaded):
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("No se pudo leer el archivo Excel. Verificá que no esté dañado.") from error
    worksheet = workbook.active
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    columns = {normalize(value): index for index, value in enumerate(header_row) if normalize(value)}
    indexes = {}
    for field, aliases in HEADER_ALIASES.items():
        indexes[field] = next((columns[alias] for alias in aliases if alias in columns), None)
    missing = [field for field in ("name", "category", "brand", "supplier", "price") if indexes[field] is None]
    if missing:
        labels = {"name": "Nombre", "category": "Categoría", "brand": "Marca", "supplier": "Proveedor", "price": "Precio"}
        return [], [{"line": 1, "message": "Faltan columnas obligatorias: " + ", ".join(labels[field] for field in missing) + "."}]
    rows, errors = [], []
    category_aliases = {"alimento": "food", "juguete": "toy", "abrigo": "clothing", "ropa": "clothing", "cama": "bed", "accesorio": "accessory", "higiene": "hygiene", "shampoo": "shampoo", "cosmetica": "shampoo", "medicamento": "medicine", "medicamento veterinario": "medicine", "pipeta": "antiparasitic", "antiparasitario": "antiparasitic", "otro": "other"}
    unit_aliases = {"unidad": "unit", "unidades": "unit", "kilo": "kg", "kilogramo": "kg", "kilogramos": "kg", "bolsa": "bag"}
    for line_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        def value_for(field):
            index = indexes[field]
            return values[index] if index is not None and index < len(values) else None
        try:
            name, brand, supplier = (as_text(value_for(field)) for field in ("name", "brand", "supplier"))
            if not name or not brand or not supplier:
                raise ValueError("Nombre, Marca y Proveedor son obligatorios.")
            row = {"sku": as_text(value_for("sku")) or None, "name": name, "commercial_name": as_text(value_for("commercial_name")), "category": choice_value(value_for("category"), Product.Category.choices, "Categoría", category_aliases), "quality": as_text(value_for("quality")), "brand": brand, "supplier": supplier, "sale_unit": choice_value(value_for("sale_unit"), Product.SaleUnit.choices, "Se vende por", unit_aliases, Product.SaleUnit.UNIT), "package_weight_kg": as_decimal(value_for("package_weight_kg"), "Peso de bolsa", default=None), "price": as_decimal(value_for("price"), "Precio", required=True), "promotional_price": as_decimal(value_for("promotional_price"), "Precio promocional", default=None), "promotion_starts_on": as_date(value_for("promotion_starts_on"), "Promoción desde"), "promotion_ends_on": as_date(value_for("promotion_ends_on"), "Promoción hasta"), "stock": as_decimal(value_for("stock"), "Stock", default=Decimal("0")), "minimum_stock": as_decimal(value_for("minimum_stock"), "Stock mínimo", default=Decimal("0")), "batch": as_text(value_for("batch")), "expires_on": as_date(value_for("expires_on"), "Vencimiento"), "requires_prescription": as_bool(value_for("requires_prescription"))}
            if row["promotional_price"] is not None and row["promotion_starts_on"] and row["promotion_ends_on"] and row["promotion_ends_on"] < row["promotion_starts_on"]:
                raise ValueError("La promoción termina antes de comenzar.")
            rows.append(row)
        except ValueError as error:
            errors.append({"line": line_number, "message": str(error)})
    return rows, errors
